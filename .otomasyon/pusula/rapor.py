# -*- coding: utf-8 -*-
"""Toplu rapor — Excel (openpyxl varsa) + her zaman CSV + özet HTML panosu."""
import os, csv, json, html, datetime
from .ayarlar import CIKTI, SEKTORLER
from .puanlama import ETIKET, sicak_mi, bizim_eksikler

BASLIKLAR = ["ID","İşletme","Sektör","Şehir","İlçe","Telefon","Site","Puan","Yorum",
             "Skor","Sıcak mı","Eksik sayısı","Luna'nın çözdüğü","Eksikler",
             "Aylık görüntülenme (şu an)","Aylık görüntülenme (hedef)","Artış %",
             "Aylık iletişim (şu an)","Aylık iletişim (hedef)","Aylık ek ciro (tahmin)"]

def _satir(r, t):
    eks = json.loads(r["eksikler"])
    return [r["id"], r["ad"], SEKTORLER.get(r["sektor"],{}).get("ad", r["sektor"] or ""),
            r["sehir"] or "", r["ilce"] or "", r["telefon"] or "", r["site"] or "",
            r["puan"] or "", r["yorum_sayisi"] or 0, r["skor"],
            "EVET" if sicak_mi(r["skor"]) else "hayır",
            len(eks), len(bizim_eksikler(eks)),
            ", ".join(ETIKET.get(k,k) for k in eks),
            (t or {}).get("mevcut_goruntulenme",""), (t or {}).get("hedef_goruntulenme",""),
            "", (t or {}).get("mevcut_iletisim",""), (t or {}).get("hedef_iletisim",""),
            (t or {}).get("ek_gelir","")]

def uret(satirlar, tahminler, ad="pusula-rapor"):
    os.makedirs(CIKTI, exist_ok=True)
    damga = datetime.date.today().strftime("%Y%m%d")
    veriler = []
    for r in satirlar:
        t = tahminler.get(r["id"])
        s = _satir(r, t)
        if t and t.get("mevcut_goruntulenme"):
            s[16] = int(round((t["hedef_goruntulenme"]/t["mevcut_goruntulenme"]-1)*100))
        veriler.append(s)

    csv_yolu = os.path.join(CIKTI, "%s-%s.csv" % (ad, damga))
    with open(csv_yolu, "w", encoding="utf-8-sig", newline="") as f:
        y = csv.writer(f, delimiter=";")
        y.writerow(BASLIKLAR); y.writerows(veriler)

    xlsx_yolu = None
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook(); ws = wb.active; ws.title = "Adaylar"
        ws.append(BASLIKLAR)
        for h in ws[1]:
            h.font = Font(bold=True, color="FFFFFF")
            h.fill = PatternFill("solid", fgColor="E8452C")
            h.alignment = Alignment(vertical="center", wrap_text=True)
        for s in veriler: ws.append(s)
        genis = {1:6,2:34,3:20,4:14,5:14,6:16,7:34,8:7,9:8,10:7,11:10,12:12,13:14,14:60,
                 15:16,16:16,17:9,18:16,19:16,20:18}
        for i,w in genis.items():
            ws.column_dimensions[ws.cell(row=1,column=i).column_letter].width = w
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        xlsx_yolu = os.path.join(CIKTI, "%s-%s.xlsx" % (ad, damga))
        wb.save(xlsx_yolu)
    except ImportError:
        pass
    return csv_yolu, xlsx_yolu

def pano(satirlar, tahminler, demolar, ad="pusula-pano"):
    """Tek bakışta liste — demolara tıklanabilir bağlantılarla."""
    os.makedirs(CIKTI, exist_ok=True)
    damga = datetime.date.today().strftime("%Y%m%d")
    tr = []
    for r in satirlar:
        eks = json.loads(r["eksikler"]); t = tahminler.get(r["id"]) or {}
        d = demolar.get(r["id"])
        bag = '<a href="%s">demo</a>' % html.escape(os.path.relpath(d, CIKTI)) if d else "—"
        tr.append("<tr class='%s'><td>%d</td><td><b>%s</b><br><small>%s · %s</small></td>"
                  "<td class='s'>%d</td><td>%d</td><td>%s</td><td>%s</td><td>%s</td></tr>" % (
            "sicak" if sicak_mi(r["skor"]) else "", r["id"], html.escape(r["ad"]),
            html.escape(SEKTORLER.get(r["sektor"],{}).get("ad", r["sektor"] or "")),
            html.escape(r["sehir"] or ""), r["skor"], len(eks),
            "{:,}".format(t.get("ek_gelir",0)).replace(",","."),
            html.escape(r["telefon"] or "—"), bag))
    kalip = """<!DOCTYPE html><html lang="tr"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Luna Pusula — Aday Panosu</title><style>
body{background:#0A0A0C;color:#EFEDE8;font-family:system-ui,-apple-system,"Segoe UI",sans-serif;margin:0;padding:34px 20px}
h1{font-size:26px;margin:0 0 6px}
p.alt{color:#8C8A84;margin:0 0 24px;font-size:14px}
table{width:100%;max-width:1180px;border-collapse:collapse;font-size:14.5px}
th,td{text-align:left;padding:11px 12px;border-bottom:1px solid rgba(239,237,232,.12)}
th{font-size:10.5px;letter-spacing:.15em;text-transform:uppercase;color:#E8452C}
small{color:#8C8A84}
td.s{font-weight:700}
tr.sicak td.s{color:#E8452C}
tr.sicak{background:rgba(232,69,44,.06)}
a{color:#E8452C}
</style></head><body>
<h1>Luna Pusula &mdash; aday panosu</h1>
<p class="alt">@ADET@ aday &middot; kırmızı satırlar sıcak adaylar (skoru düşük, eksiği çok) &middot; @TARIH@</p>
<div class='tablo-kaydir'><table><tr><th>#</th><th>İşletme</th><th>Skor</th><th>Eksik</th><th>Tahmini aylık ek ciro</th><th>Telefon</th><th>Demo</th></tr>
@SATIR@</table></div></body></html>"""
    g = (kalip.replace("@ADET@", str(len(satirlar)))
              .replace("@TARIH@", datetime.date.today().strftime("%d.%m.%Y"))
              .replace("@SATIR@", "".join(tr)))
    yol = os.path.join(CIKTI, "%s-%s.html" % (ad, damga))
    open(yol,"w",encoding="utf-8").write(g)
    return yol
